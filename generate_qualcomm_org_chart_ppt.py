"""Generate a consulting-style PowerPoint org chart for Qualcomm.

The deck uses public leadership information and a clean, McKinsey-inspired
visual system without relying on proprietary template assets.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUTFILE = Path("qualcomm_org_chart_mckinsey_style.pptx")

SLIDE_W = 13.333
SLIDE_H = 7.5

NAVY = RGBColor(31, 78, 121)
BLUE = RGBColor(0, 102, 204)
LIGHT_BLUE = RGBColor(226, 239, 255)
PALE_BLUE = RGBColor(242, 247, 253)
DARK = RGBColor(30, 30, 30)
MID = RGBColor(95, 95, 95)
LIGHT = RGBColor(222, 226, 231)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0, 0, 0)


EXECUTIVES = [
    {
        "name": "Akash Palkhiwala",
        "title": "EVP, CFO & COO",
        "domain": "Finance, operations, GTM, IT",
    },
    {
        "name": "Alexander H. Rogers",
        "title": "President, QTL & Global Affairs",
        "domain": "Licensing, policy, global affairs",
    },
    {
        "name": "Alex Katouzian",
        "title": "EVP & Group GM, MCX",
        "domain": "Mobile, compute and XR platforms",
    },
    {
        "name": "Dr. Baaziz Achour",
        "title": "EVP & CTO, QTI",
        "domain": "Technology roadmap, R&D, engineering",
    },
    {
        "name": "Colin Ryan",
        "title": "EVP, Chief Strategy & Corp. Dev. Officer",
        "domain": "Enterprise strategy and M&A",
    },
    {
        "name": "Ann Chaplin",
        "title": "EVP, General Counsel & Corporate Secretary",
        "domain": "Legal, governance, compliance",
    },
    {
        "name": "Heather Ace",
        "title": "EVP, Human Resources & CHRO",
        "domain": "People, talent, culture",
    },
    {
        "name": "Don McGuire",
        "title": "EVP & Chief Marketing Officer",
        "domain": "Brand, communications, marketing",
    },
    {
        "name": "Thomas Ta",
        "title": "Chief Artificial Intelligence Officer",
        "domain": "AI strategy and enablement",
    },
]


PILLARS = [
    (
        "Business leadership",
        [
            ("Alexander H. Rogers", "QTL and global affairs"),
            ("Alex Katouzian", "Mobile, compute and XR business unit"),
        ],
    ),
    (
        "Technology & product",
        [
            ("Dr. Baaziz Achour", "R&D, engineering and technology strategy"),
            ("Thomas Ta", "Company-wide AI agenda"),
        ],
    ),
    (
        "Corporate center",
        [
            ("Akash Palkhiwala", "Finance, operations, GTM, IT"),
            ("Colin Ryan", "Strategy and corporate development"),
        ],
    ),
    (
        "Enabling functions",
        [
            ("Ann Chaplin", "Legal and corporate secretary"),
            ("Heather Ace", "Human resources"),
            ("Don McGuire", "Marketing and communications"),
        ],
    ),
]


SOURCES = [
    "Qualcomm public leadership pages for Cristiano Amon, Akash Palkhiwala, Ann Chaplin, Colin Ryan, Alex Katouzian, Don McGuire and Baaziz Achour.",
    "Qualcomm press releases on Akash Palkhiwala's CFO/COO appointment, Colin Ryan's strategy appointment and Baaziz Achour's CTO succession.",
    "The Org Qualcomm leadership-team listing used as a cross-check where public pages were JavaScript-limited in automated retrieval.",
]


def inches(value: float):
    return Inches(value)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    font_size: int = 14,
    bold: bool = False,
    color: RGBColor = DARK,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    font: str = "Arial",
):
    box = slide.shapes.add_textbox(inches(x), inches(y), inches(w), inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_header(slide, title: str, eyebrow: str = "QUALCOMM | public leadership view"):
    add_text(slide, title, 0.55, 0.28, 9.7, 0.48, font_size=18, bold=True)
    add_text(slide, eyebrow, 10.2, 0.32, 2.55, 0.32, font_size=8, color=MID, align=PP_ALIGN.RIGHT)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, inches(0.55), inches(0.9), inches(12.2), inches(0.025))
    line.fill.solid()
    line.fill.fore_color.rgb = BLUE
    line.line.fill.background()


def add_footer(slide, page: int):
    add_text(
        slide,
        "Source: public Qualcomm leadership pages, press releases and SEC/company materials; org view prepared May 8, 2026",
        0.55,
        7.05,
        10.7,
        0.18,
        font_size=6,
        color=MID,
    )
    add_text(slide, str(page), 12.4, 7.03, 0.35, 0.2, font_size=7, color=MID, align=PP_ALIGN.RIGHT)


def add_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: RGBColor,
    line: RGBColor,
    radius: MSO_SHAPE = MSO_SHAPE.ROUNDED_RECTANGLE,
):
    shape = slide.shapes.add_shape(radius, inches(x), inches(y), inches(w), inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.0)
    return shape


def add_person_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    name: str,
    title: str,
    domain: str,
    *,
    accent: RGBColor = BLUE,
    fill: RGBColor = WHITE,
):
    shape = add_box(slide, x, y, w, h, fill=fill, line=LIGHT)
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, inches(x), inches(y), inches(0.08), inches(h))
    strip.fill.solid()
    strip.fill.fore_color.rgb = accent
    strip.line.fill.background()
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = inches(0.15)
    frame.margin_right = inches(0.08)
    frame.margin_top = inches(0.07)
    frame.margin_bottom = inches(0.05)

    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = name
    r.font.name = "Arial"
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = DARK

    p = frame.add_paragraph()
    p.space_before = Pt(1)
    r = p.add_run()
    r.text = title
    r.font.name = "Arial"
    r.font.size = Pt(6.7)
    r.font.color.rgb = NAVY

    p = frame.add_paragraph()
    p.space_before = Pt(1)
    r = p.add_run()
    r.text = domain
    r.font.name = "Arial"
    r.font.size = Pt(5.8)
    r.font.color.rgb = MID
    return shape


def connect(slide, x1: float, y1: float, x2: float, y2: float, *, color: RGBColor = LIGHT, width: float = 1.1):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, inches(x1), inches(y1), inches(x2), inches(y2)
    )
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    return connector


def set_deck_properties(prs: Presentation):
    props = prs.core_properties
    props.title = "Qualcomm executive organization overview"
    props.subject = "Public executive-level org chart in consulting-style PowerPoint format"
    props.author = "Cursor Agent"
    props.keywords = "Qualcomm, org chart, executive leadership, PowerPoint"


def build_cover(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE

    add_text(slide, "Qualcomm Incorporated", 0.72, 1.58, 6.2, 0.45, font_size=19, color=NAVY, bold=True)
    add_text(slide, "Executive organization overview", 0.72, 2.13, 7.6, 0.55, font_size=30, bold=True, color=BLACK)
    add_text(
        slide,
        "Public leadership view | McKinsey-style presentation format",
        0.72,
        2.84,
        7.5,
        0.35,
        font_size=12,
        color=MID,
    )
    add_text(slide, "Prepared May 8, 2026", 0.72, 5.82, 3.0, 0.24, font_size=9, color=MID)

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, inches(10.4), inches(0), inches(2.92), inches(SLIDE_H))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    add_text(slide, "QCOM", 10.85, 0.72, 1.85, 0.42, font_size=24, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)
    add_text(
        slide,
        "Executive leadership\nstructure based on\npublic sources",
        10.86,
        5.78,
        1.85,
        0.84,
        font_size=10,
        color=WHITE,
        align=PP_ALIGN.RIGHT,
    )

    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, inches(0.72), inches(3.42), inches(3.1), inches(0.05))
    accent.fill.solid()
    accent.fill.fore_color.rgb = BLUE
    accent.line.fill.background()
    return slide


def build_org_chart(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Qualcomm executive organization chart")

    board = add_box(slide, 4.78, 1.14, 3.7, 0.48, fill=PALE_BLUE, line=LIGHT)
    board.text_frame.text = "Board of Directors\nChair: Mark D. McLaughlin"
    for paragraph in board.text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8 if paragraph.text.startswith("Board") else 6.5)
            run.font.bold = paragraph.text.startswith("Board")
            run.font.color.rgb = DARK

    ceo = add_box(slide, 4.52, 1.98, 4.24, 0.74, fill=NAVY, line=NAVY)
    ceo.text_frame.clear()
    ceo.text_frame.margin_top = inches(0.08)
    ceo.text_frame.margin_bottom = inches(0.05)
    ceo.text_frame.margin_left = inches(0.12)
    ceo.text_frame.margin_right = inches(0.12)
    p = ceo.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Cristiano R. Amon"
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = WHITE
    p = ceo.text_frame.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "President & Chief Executive Officer"
    r.font.name = "Arial"
    r.font.size = Pt(8.5)
    r.font.color.rgb = WHITE

    connect(slide, 6.63, 1.62, 6.63, 1.98, color=LIGHT)

    add_text(
        slide,
        "Direct executive reports / leadership roles shown from public sources",
        0.62,
        2.96,
        4.0,
        0.2,
        font_size=7,
        color=MID,
    )

    trunk_y = 3.18
    connect(slide, 6.63, 2.72, 6.63, trunk_y, color=LIGHT)
    connect(slide, 1.35, trunk_y, 11.97, trunk_y, color=LIGHT)

    box_w = 2.22
    box_h = 0.82
    gap = 0.20
    left = 0.62
    row_ys = [3.46, 4.62]
    for idx, person in enumerate(EXECUTIVES):
        row = 0 if idx < 5 else 1
        col = idx if idx < 5 else idx - 5
        x = left + col * (box_w + gap)
        y = row_ys[row]
        add_person_box(slide, x, y, box_w, box_h, person["name"], person["title"], person["domain"])
        connect(slide, x + box_w / 2, trunk_y, x + box_w / 2, y, color=LIGHT)

    note = add_box(slide, 0.62, 6.03, 12.08, 0.42, fill=PALE_BLUE, line=LIGHT, radius=MSO_SHAPE.RECTANGLE)
    note.text_frame.clear()
    note.text_frame.margin_left = inches(0.12)
    note.text_frame.margin_right = inches(0.12)
    note.text_frame.margin_top = inches(0.06)
    p = note.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "Note: chart reflects an executive-level public view; internal reporting lines, full staff organizations and interim changes are not represented."
    r.font.name = "Arial"
    r.font.size = Pt(7)
    r.font.color.rgb = MID

    add_footer(slide, 2)
    return slide


def build_function_view(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Leadership roles cluster around four executive agendas")

    add_text(
        slide,
        "The public leadership roster maps to business-line ownership, technology leadership, corporate-center controls and enabling functions.",
        0.58,
        1.08,
        11.9,
        0.3,
        font_size=11,
        color=MID,
    )

    card_w = 2.92
    gap = 0.18
    x0 = 0.58
    for idx, (heading, entries) in enumerate(PILLARS):
        x = x0 + idx * (card_w + gap)
        add_box(slide, x, 1.65, card_w, 4.72, fill=WHITE, line=LIGHT, radius=MSO_SHAPE.RECTANGLE)
        top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, inches(x), inches(1.65), inches(card_w), inches(0.52))
        top.fill.solid()
        top.fill.fore_color.rgb = NAVY if idx in (0, 1) else BLUE
        top.line.fill.background()
        add_text(slide, heading, x + 0.15, 1.79, card_w - 0.3, 0.18, font_size=9.5, bold=True, color=WHITE)

        y = 2.43
        for name, description in entries:
            bullet = slide.shapes.add_shape(MSO_SHAPE.OVAL, inches(x + 0.18), inches(y + 0.02), inches(0.08), inches(0.08))
            bullet.fill.solid()
            bullet.fill.fore_color.rgb = BLUE
            bullet.line.fill.background()
            add_text(slide, name, x + 0.34, y, card_w - 0.55, 0.18, font_size=8.2, bold=True, color=DARK)
            add_text(slide, description, x + 0.34, y + 0.22, card_w - 0.55, 0.35, font_size=7.1, color=MID)
            y += 0.82

    add_footer(slide, 3)
    return slide


def build_sources(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Sources and assumptions")

    add_text(slide, "Source base", 0.72, 1.25, 2.4, 0.28, font_size=14, bold=True)
    y = 1.75
    for item in SOURCES:
        add_text(slide, u"\u2022", 0.78, y, 0.12, 0.18, font_size=11, color=BLUE)
        add_text(slide, item, 1.02, y, 10.6, 0.42, font_size=9, color=DARK)
        y += 0.58

    add_text(slide, "Interpretation choices", 0.72, 4.1, 3.1, 0.28, font_size=14, bold=True)
    assumptions = [
        "Visual hierarchy places Cristiano R. Amon as President & CEO, with public executive leadership roles arrayed underneath.",
        "Board chair is shown as oversight context, not as part of day-to-day management.",
        "The file uses a consulting-style layout because no official McKinsey template asset was present in the repository.",
    ]
    y = 4.6
    for item in assumptions:
        add_text(slide, u"\u2022", 0.78, y, 0.12, 0.18, font_size=11, color=BLUE)
        add_text(slide, item, 1.02, y, 10.6, 0.36, font_size=9, color=DARK)
        y += 0.52

    add_footer(slide, 4)
    return slide


def main():
    prs = Presentation()
    prs.slide_width = inches(SLIDE_W)
    prs.slide_height = inches(SLIDE_H)
    set_deck_properties(prs)
    build_cover(prs)
    build_org_chart(prs)
    build_function_view(prs)
    build_sources(prs)
    prs.save(OUTFILE)
    print(f"Wrote {OUTFILE}")


if __name__ == "__main__":
    main()
